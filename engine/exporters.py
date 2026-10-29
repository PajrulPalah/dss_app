"""
Ekspor hasil DSS ke format Excel (dengan rumus hidup), CSV, dan Markdown.
Tidak mengimpor streamlit.
"""
from __future__ import annotations

import io
from typing import Optional
import pandas as pd
import numpy as np

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from engine.models import Problem, SAWResult, WPResult
from config import BENEFIT


# --- Warna Excel ---
FILL_INPUT = PatternFill("solid", fgColor="EAF4FB") if HAS_OPENPYXL else None
FILL_FORMULA = PatternFill("solid", fgColor="FFFFFF") if HAS_OPENPYXL else None
FILL_HEADER = PatternFill("solid", fgColor="0F4C5C") if HAS_OPENPYXL else None
FILL_ACCENT = PatternFill("solid", fgColor="D9EAEC") if HAS_OPENPYXL else None
FONT_HEADER = Font(bold=True, color="FFFFFF") if HAS_OPENPYXL else None
FONT_FORMULA = Font(color="000000") if HAS_OPENPYXL else None
FONT_INPUT = Font(color="1155CC", bold=False) if HAS_OPENPYXL else None


def _auto_width(ws) -> None:
    """Sesuaikan lebar kolom secara otomatis."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _header_style(cell, teks: str) -> None:
    cell.value = teks
    cell.fill = FILL_HEADER
    cell.font = FONT_HEADER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _tulis_sheet_kriteria(wb, problem: Problem) -> None:
    """Tulis sheet Kriteria dengan rumus hidup."""
    ws = wb.create_sheet("Kriteria")
    ws.merge_cells("A1:G1")
    ws["A1"] = "Kriteria & Bobot"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["Kode", "Nama", "Tipe", "Bobot Input", "Bobot Normalisasi (w)", "Pangkat WP", "Alasan"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=j)
        _header_style(c, h)

    n = len(problem.kriteria)
    sum_row = 3 + n + 1  # baris Total

    for i, k in enumerate(problem.kriteria):
        row = 4 + i
        ws.cell(row=row, column=1, value=k.kode).fill = FILL_ACCENT
        ws.cell(row=row, column=2, value=k.nama).fill = FILL_ACCENT
        ws.cell(row=row, column=3, value=k.tipe).fill = FILL_INPUT
        c_bobot = ws.cell(row=row, column=4, value=k.bobot_input)
        c_bobot.fill = FILL_INPUT
        c_bobot.font = FONT_INPUT
        # Bobot normalisasi = D{row}/SUM(D4:D{4+n-1})
        d_start, d_end = 4, 3 + n
        ws.cell(row=row, column=5, value=f"=D{row}/SUM($D${d_start}:$D${d_end})").font = FONT_FORMULA
        # Pangkat WP = IF(C{row}="Cost", -E{row}, E{row})
        ws.cell(row=row, column=6, value=f'=IF(C{row}="Cost",-E{row},E{row})').font = FONT_FORMULA
        ws.cell(row=row, column=7, value=k.alasan).fill = FILL_ACCENT

    ws.cell(row=sum_row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=sum_row, column=4, value=f"=SUM(D4:D{3+n})").font = Font(bold=True)
    ws.cell(row=sum_row, column=5, value=f"=SUM(E4:E{3+n})").font = Font(bold=True)

    ws.cell(row=sum_row + 2, column=1, value="Nilai pengganti data ≤ 0")
    ws.cell(row=sum_row + 2, column=4, value=0.01).fill = FILL_INPUT

    _auto_width(ws)


def _tulis_sheet_data(wb, problem: Problem) -> None:
    """Tulis sheet Data."""
    ws = wb.create_sheet("Data")
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Data Alternatif — {problem.nama_kasus}"
    ws["A1"].font = Font(bold=True, size=12)

    headers = ["Alternatif"] + [f"{k.kode} {k.nama}" for k in problem.kriteria]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=j)
        _header_style(c, h)

    for i, alt in enumerate(problem.alternatif):
        row = 4 + i
        ws.cell(row=row, column=1, value=alt).fill = FILL_ACCENT
        for j, k in enumerate(problem.kriteria):
            val = problem.matriks.loc[alt, k.kode]
            c = ws.cell(row=row, column=2 + j, value=float(val))
            c.fill = FILL_INPUT
            c.font = FONT_INPUT

    _auto_width(ws)


def _tulis_sheet_saw(wb, hasil: SAWResult) -> None:
    """Tulis sheet Perhitungan_SAW."""
    ws = wb.create_sheet("Perhitungan_SAW")
    problem = hasil.problem
    n_alt = len(problem.alternatif)
    n_crit = len(problem.kriteria)
    kode = problem.kode_kriteria

    ws["A1"] = "Perhitungan SAW (Simple Additive Weighting)"
    ws["A1"].font = Font(bold=True, size=12)

    # Bobot
    ws["A3"] = "Bobot Ternormalisasi (w)"
    ws["A3"].font = Font(bold=True)
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=4, column=2 + j, value=k.kode).font = Font(bold=True)
        ws.cell(row=5, column=2 + j, value=round(float(hasil.bobot[j]), 6))

    # Pembagi normalisasi
    ws["A7"] = "Pembagi Normalisasi"
    ws["A7"].font = Font(bold=True)
    ws["B7"] = "Tipe"
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=8, column=2 + j, value=k.kode).font = Font(bold=True)
        ws.cell(row=9, column=1, value="Pembagi (max/min)")
        ws.cell(row=9, column=2 + j, value=round(float(hasil.pembagi[k.kode]), 6))

    # Matriks R
    row_start_r = 12
    ws.cell(row=row_start_r, column=1, value="Matriks R (Ternormalisasi)").font = Font(bold=True)
    ws.cell(row=row_start_r + 1, column=1, value="Alternatif").font = Font(bold=True)
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=row_start_r + 1, column=2 + j, value=k.kode).font = Font(bold=True)
    for i, alt in enumerate(problem.alternatif):
        ws.cell(row=row_start_r + 2 + i, column=1, value=alt)
        for j, k in enumerate(problem.kriteria):
            ws.cell(row=row_start_r + 2 + i, column=2 + j, value=round(float(hasil.matriks_r.loc[alt, k.kode]), 6))

    # Matriks w*r dan V
    row_start_v = row_start_r + 2 + n_alt + 2
    ws.cell(row=row_start_v, column=1, value="Matriks w×r dan Skor V").font = Font(bold=True)
    ws.cell(row=row_start_v + 1, column=1, value="Alternatif").font = Font(bold=True)
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=row_start_v + 1, column=2 + j, value=f"w×{k.kode}").font = Font(bold=True)
    ws.cell(row=row_start_v + 1, column=2 + n_crit, value="V (Skor)").font = Font(bold=True)
    ws.cell(row=row_start_v + 1, column=3 + n_crit, value="Peringkat").font = Font(bold=True)

    for i, alt in enumerate(problem.alternatif):
        r = row_start_v + 2 + i
        ws.cell(row=r, column=1, value=alt)
        # Kolom w*r
        wr_cols = [get_column_letter(2 + j) for j in range(n_crit)]
        for j, k in enumerate(problem.kriteria):
            ws.cell(row=r, column=2 + j, value=round(float(hasil.matriks_wr.loc[alt, k.kode]), 6))
        # V = SUM(w*r kolom)
        v_col = get_column_letter(2 + n_crit)
        col_start = get_column_letter(2)
        col_end = get_column_letter(1 + n_crit)
        ws.cell(row=r, column=2 + n_crit, value=f"=SUM({col_start}{r}:{col_end}{r})")
        ws.cell(row=r, column=3 + n_crit, value=int(hasil.peringkat.loc[alt]))

    _auto_width(ws)


def _tulis_sheet_wp(wb, hasil: WPResult) -> None:
    """Tulis sheet Perhitungan_WP."""
    ws = wb.create_sheet("Perhitungan_WP")
    problem = hasil.problem
    n_alt = len(problem.alternatif)
    n_crit = len(problem.kriteria)

    ws["A1"] = "Perhitungan WP (Weighted Product)"
    ws["A1"].font = Font(bold=True, size=12)

    # Pangkat
    ws["A3"] = "Pangkat per Kriteria (+w=Benefit, -w=Cost)"
    ws["A3"].font = Font(bold=True)
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=4, column=2 + j, value=k.kode).font = Font(bold=True)
        ws.cell(row=5, column=2 + j, value=round(float(hasil.pangkat[j]), 6))

    # Matriks yang disesuaikan
    row_adj = 8
    ws.cell(row=row_adj, column=1, value="Matriks Disesuaikan (nilai ≤ 0 diganti)").font = Font(bold=True)
    ws.cell(row=row_adj + 1, column=1, value="Alternatif").font = Font(bold=True)
    for j, k in enumerate(problem.kriteria):
        ws.cell(row=row_adj + 1, column=2 + j, value=k.kode).font = Font(bold=True)
    for i, alt in enumerate(problem.alternatif):
        ws.cell(row=row_adj + 2 + i, column=1, value=alt)
        for j, k in enumerate(problem.kriteria):
            ws.cell(row=row_adj + 2 + i, column=2 + j, value=round(float(hasil.matriks_adjusted.loc[alt, k.kode]), 6))

    # S dan V
    row_sv = row_adj + 2 + n_alt + 2
    ws.cell(row=row_sv, column=1, value="Vektor S dan V").font = Font(bold=True)
    ws.cell(row=row_sv + 1, column=1, value="Alternatif").font = Font(bold=True)
    ws.cell(row=row_sv + 1, column=2, value="S").font = Font(bold=True)
    ws.cell(row=row_sv + 1, column=3, value="V = S/ΣS").font = Font(bold=True)
    ws.cell(row=row_sv + 1, column=4, value="Peringkat").font = Font(bold=True)

    s_rows = [row_sv + 2 + i for i in range(n_alt)]
    for i, alt in enumerate(problem.alternatif):
        r = s_rows[i]
        ws.cell(row=r, column=1, value=alt)
        ws.cell(row=r, column=2, value=round(float(hasil.s.loc[alt]), 6))
        # V = B{r} / SUM(B{s_rows[0]}:B{s_rows[-1]})
        s_range = f"B{s_rows[0]}:B{s_rows[-1]}"
        ws.cell(row=r, column=3, value=f"=B{r}/SUM({s_range})")
        ws.cell(row=r, column=4, value=int(hasil.peringkat.loc[alt]))

    # Total S
    sum_row = s_rows[-1] + 1
    ws.cell(row=sum_row, column=1, value="Σ S").font = Font(bold=True)
    ws.cell(row=sum_row, column=2, value=f"=SUM(B{s_rows[0]}:B{s_rows[-1]})").font = Font(bold=True)
    ws.cell(row=sum_row, column=3, value=f"=SUM(C{s_rows[0]}:C{s_rows[-1]})").font = Font(bold=True)

    _auto_width(ws)


def _tulis_sheet_ranking(wb, hasil_saw: Optional[SAWResult], hasil_wp: Optional[WPResult]) -> None:
    """Tulis sheet Hasil Ranking."""
    ws = wb.create_sheet("Hasil Ranking")
    ws["A1"] = "Hasil Ranking MADM"
    ws["A1"].font = Font(bold=True, size=13)

    headers = ["Alternatif"]
    if hasil_saw:
        headers += ["V SAW", "Peringkat SAW"]
    if hasil_wp:
        headers += ["V WP", "Peringkat WP"]

    for j, h in enumerate(headers, 1):
        _header_style(ws.cell(row=3, column=j), h)

    # Tentukan alternatif dari salah satu hasil
    problem = (hasil_saw or hasil_wp).problem
    for i, alt in enumerate(problem.alternatif):
        r = 4 + i
        ws.cell(row=r, column=1, value=alt)
        col = 2
        if hasil_saw:
            ws.cell(row=r, column=col, value=round(float(hasil_saw.v.loc[alt]), 6))
            ws.cell(row=r, column=col + 1, value=int(hasil_saw.peringkat.loc[alt]))
            col += 2
        if hasil_wp:
            ws.cell(row=r, column=col, value=round(float(hasil_wp.v.loc[alt]), 6))
            ws.cell(row=r, column=col + 1, value=int(hasil_wp.peringkat.loc[alt]))

    _auto_width(ws)


def _tulis_sheet_panduan(wb) -> None:
    """Tulis sheet Panduan."""
    ws = wb.create_sheet("Panduan")
    ws["A1"] = "Panduan Penggunaan DSS Calculator Engine"
    ws["A1"].font = Font(bold=True, size=14)
    panduan = [
        ("SAW (Simple Additive Weighting)", "Normalisasi matriks keputusan → hitung V = Σ w_j × r_ij → ranking."),
        ("WP (Weighted Product)", "Hitung S_i = ∏ x_ij^pangkat_j → V_i = S_i / Σ S → ranking."),
        ("Bobot", "Input bebas (mis. 20/30/50); selalu dinormalisasi agar Σw = 1."),
        ("Nilai ≤ 0 (WP)", "WP memerlukan nilai positif; strategi: ganti, geser, atau tolak."),
        ("Peringkat", "1 = terbaik. Tie menggunakan metode 'min' (peringkat sama)."),
    ]
    for i, (topik, penjelasan) in enumerate(panduan, 3):
        ws.cell(row=i, column=1, value=topik).font = Font(bold=True)
        ws.cell(row=i, column=2, value=penjelasan)
    _auto_width(ws)


def ekspor_excel(
    problem: Problem,
    hasil_saw: Optional[SAWResult] = None,
    hasil_wp: Optional[WPResult] = None,
) -> bytes:
    """
    Ekspor hasil ke file Excel (.xlsx) dengan rumus hidup.

    Parameters
    ----------
    problem : Problem
        Masalah MADM.
    hasil_saw : SAWResult, optional
        Hasil SAW.
    hasil_wp : WPResult, optional
        Hasil WP.

    Returns
    -------
    bytes
        Konten file Excel.
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl tidak terinstall. Jalankan: pip install openpyxl")

    wb = openpyxl.Workbook()
    # Hapus sheet default
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    _tulis_sheet_kriteria(wb, problem)
    _tulis_sheet_data(wb, problem)
    if hasil_saw:
        _tulis_sheet_saw(wb, hasil_saw)
    if hasil_wp:
        _tulis_sheet_wp(wb, hasil_wp)
    _tulis_sheet_ranking(wb, hasil_saw, hasil_wp)
    _tulis_sheet_panduan(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ekspor_csv_ranking(
    problem: Problem,
    hasil_saw: Optional[SAWResult] = None,
    hasil_wp: Optional[WPResult] = None,
) -> str:
    """
    Ekspor tabel ranking ke format CSV.

    Returns
    -------
    str
        Konten CSV.
    """
    rows = []
    for alt in problem.alternatif:
        row = {"Alternatif": alt}
        if hasil_saw:
            row["V_SAW"] = round(float(hasil_saw.v.loc[alt]), 6)
            row["Peringkat_SAW"] = int(hasil_saw.peringkat.loc[alt])
        if hasil_wp:
            row["V_WP"] = round(float(hasil_wp.v.loc[alt]), 6)
            row["Peringkat_WP"] = int(hasil_wp.peringkat.loc[alt])
        rows.append(row)
    df = pd.DataFrame(rows)
    return df.to_csv(index=False)


def ekspor_markdown(
    problem: Problem,
    hasil_saw: Optional[SAWResult] = None,
    hasil_wp: Optional[WPResult] = None,
) -> str:
    """
    Ekspor laporan ringkas ke format Markdown.

    Returns
    -------
    str
        Konten Markdown.
    """
    lines = [
        f"# Laporan DSS: {problem.nama_kasus}",
        "",
        f"**Pengambil Keputusan:** {problem.pengambil_keputusan}",
        f"**Tujuan:** {problem.tujuan}",
        f"**Sumber Data:** {problem.sumber_data}",
        "",
        "## Kriteria",
        "",
        "| Kode | Nama | Tipe | Bobot |",
        "|------|------|------|-------|",
    ]
    bobot = problem.bobot_ternormalisasi
    for i, k in enumerate(problem.kriteria):
        lines.append(f"| {k.kode} | {k.nama} | {k.tipe} | {bobot[i]:.4f} |")

    if hasil_saw:
        lines += [
            "",
            "## Hasil SAW",
            "",
            "| Alternatif | V | Peringkat |",
            "|-----------|---|-----------|",
        ]
        for alt in problem.alternatif:
            lines.append(f"| {alt} | {hasil_saw.v.loc[alt]:.4f} | {int(hasil_saw.peringkat.loc[alt])} |")

    if hasil_wp:
        lines += [
            "",
            "## Hasil WP",
            "",
            "| Alternatif | S | V | Peringkat |",
            "|-----------|---|---|-----------|",
        ]
        for alt in problem.alternatif:
            lines.append(f"| {alt} | {hasil_wp.s.loc[alt]:.4f} | {hasil_wp.v.loc[alt]:.4f} | {int(hasil_wp.peringkat.loc[alt])} |")

    # Kesimpulan otomatis
    lines += ["", "## Kesimpulan"]
    if hasil_saw and hasil_wp:
        best_saw = hasil_saw.peringkat.idxmin()
        best_wp = hasil_wp.peringkat.idxmin()
        if best_saw == best_wp:
            lines.append(f"Kedua metode (SAW dan WP) merekomendasikan **{best_saw}** sebagai alternatif terbaik.")
        else:
            lines.append(
                f"SAW merekomendasikan **{best_saw}**, sedangkan WP merekomendasikan **{best_wp}**. "
                f"Perbedaan ini kemungkinan disebabkan oleh sensitivitas WP terhadap nilai ekstrem."
            )
    elif hasil_saw:
        best = hasil_saw.peringkat.idxmin()
        lines.append(f"SAW merekomendasikan **{best}** sebagai alternatif terbaik.")
    elif hasil_wp:
        best = hasil_wp.peringkat.idxmin()
        lines.append(f"WP merekomendasikan **{best}** sebagai alternatif terbaik.")

    return "\n".join(lines)
